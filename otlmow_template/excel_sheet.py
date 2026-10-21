"""The writer that fills the sheets of an Excel template.

It adds the validations of the columns, the row with the definition of every
attribute and the row that marks the deprecated attributes.
"""

from asyncio import sleep
from pathlib import Path

from openpyxl.reader.excel import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import ColumnDimension, DimensionHolder
from openpyxl.worksheet.worksheet import Worksheet
from otlmow_converter.DotnotationHelper import DotnotationHelper
from otlmow_model.OtlmowModel.BaseClasses.OTLObject import OTLObject
from otlmow_template.attribute_info import AttributeInfo
from otlmow_template.constants import AGENT_TYPE_URI, MAX_VALIDATION_ROW
from otlmow_template.excel_choice_list import ExcelChoiceListMixin
from otlmow_template.excel_context import ExcelSheetContext


class ExcelSheetMixin(ExcelChoiceListMixin):
    """The class methods that fill the sheets of an Excel template."""

    @classmethod
    def alter_excel_template(
        cls,
        instances: list[OTLObject],
        file_path: Path,
        add_attribute_info: bool,
        generate_choice_list: bool,
        dummy_data_rows: int,
        add_deprecated: bool,
        abbreviate_excel_sheettitles: bool,
    ):
        wb = load_workbook(file_path)
        wb.create_sheet('Keuzelijsten')

        choice_list_dict: dict[str, str] = {}
        for sheet in wb:
            if sheet.title == 'Keuzelijsten':
                break

            cls.alter_excel_sheet(
                add_attribute_info=add_attribute_info,
                choice_list_dict=choice_list_dict,
                generate_choice_list=generate_choice_list,
                dummy_data_rows=dummy_data_rows,
                instances=instances,
                sheet=sheet,
                add_deprecated=add_deprecated,
                workbook=wb,
                abbreviate_excel_sheettitle=abbreviate_excel_sheettitles,
            )

        wb.save(file_path)
        wb.close()

    @classmethod
    async def alter_excel_template_async(
        cls,
        instances: list[OTLObject],
        file_path: Path,
        add_attribute_info: bool,
        generate_choice_list: bool,
        dummy_data_rows: int,
        add_deprecated: bool,
        abbreviate_excel_sheettitles: bool,
    ):
        wb = load_workbook(file_path)
        wb.create_sheet('Keuzelijsten')

        choice_list_dict: dict[str, str] = {}
        for sheet in wb:
            if sheet.title == 'Keuzelijsten':
                break

            cls.alter_excel_sheet(
                add_attribute_info=add_attribute_info,
                choice_list_dict=choice_list_dict,
                generate_choice_list=generate_choice_list,
                dummy_data_rows=dummy_data_rows,
                instances=instances,
                sheet=sheet,
                add_deprecated=add_deprecated,
                workbook=wb,
                abbreviate_excel_sheettitle=abbreviate_excel_sheettitles,
            )
            await sleep(0)

        wb.save(file_path)
        wb.close()

    @classmethod
    def alter_excel_sheet(
        cls,
        add_attribute_info: bool,
        choice_list_dict: dict[str, str],
        generate_choice_list: bool,
        instances: list[OTLObject],
        sheet: Worksheet,
        add_deprecated: bool,
        workbook: Workbook,
        dummy_data_rows: int,
        abbreviate_excel_sheettitle: bool = True,
    ):
        type_uri = cls.get_uri_from_sheet_name(sheet.title)
        instance = cls.get_instance_for_type_uri(instances=instances, type_uri=type_uri)

        boolean_validation = DataValidation(
            type='list', formula1='"TRUE,FALSE,-"', allow_blank=True
        )
        sheet.add_data_validation(boolean_validation)

        context = ExcelSheetContext(
            sheet=sheet,
            workbook=workbook,
            instance=instance,
            type_uri=type_uri,
            first_data_row=cls.get_first_data_row(add_attribute_info=add_attribute_info),
            boolean_validation=boolean_validation,
            generate_choice_list=generate_choice_list,
            choice_list_dict=choice_list_dict,
            info=AttributeInfo(
                add_attribute_info=add_attribute_info, add_deprecated=add_deprecated
            ),
        )

        for header_cell in next(sheet.iter_rows(min_row=1, max_row=1)):
            cls.process_header_cell(context=context, header_cell=header_cell)

        cls.format_string_columns(context=context)
        cls.delete_dummy_rows(context=context, instances=instances, dummy_data_rows=dummy_data_rows)
        cls.add_deprecated_row(context=context)
        cls.add_attribute_info_row(context=context)

        cls.set_fixed_column_width(sheet=sheet, width=25)

        if abbreviate_excel_sheettitle:
            cls.abbreviate_excel_sheettitle(sheet=sheet)

    @classmethod
    def process_header_cell(cls, context: ExcelSheetContext, header_cell) -> None:
        """Adds the validations and the information of one header cell to the sheet."""
        header = header_cell.value
        if not header:
            return
        if header == 'typeURI':
            cls.process_type_uri_header(context=context, header_cell=header_cell)
            return
        if cls.is_agent_asset_id_header(context=context, header=header):
            return
        cls.process_attribute_header(context=context, header_cell=header_cell, header=header)

    @classmethod
    def is_agent_asset_id_header(cls, context: ExcelSheetContext, header: str) -> bool:
        """Returns whether a header is the asset id of the agent sheet."""
        return context.type_uri == AGENT_TYPE_URI and header.startswith('assetId.')

    @classmethod
    def process_type_uri_header(cls, context: ExcelSheetContext, header_cell) -> None:
        """Adds the validation that fixes the type of the objects of a sheet."""
        data_validation = DataValidation(
            type='list', formula1=f'"{context.type_uri}"', allow_blank=True
        )
        context.sheet.add_data_validation(data_validation)
        data_validation.add(
            cls.get_validation_range(header_cell=header_cell, first_data_row=context.first_data_row)
        )
        context.info.add_type_uri()

    @classmethod
    def process_attribute_header(cls, context: ExcelSheetContext, header_cell, header: str) -> None:
        """Adds the validations and the information of one attribute column to the sheet."""
        attribute = DotnotationHelper.get_attribute_by_dotnotation(context.instance, header)

        context.info.add_attribute(instance=context.instance, header=header, attribute=attribute)
        cls.add_choice_list_validation(
            context=context, attribute=attribute, header_cell=header_cell
        )

        if attribute.field.native_type is str:
            context.string_columns.append(header_cell.column)

    @classmethod
    def format_string_columns(cls, context: ExcelSheetContext) -> None:
        """Formats the string columns of a sheet as text, so Excel keeps the leading zeros."""
        for col_idx in context.string_columns:
            for row_idx in range(context.first_data_row, MAX_VALIDATION_ROW + 1):
                context.sheet.cell(row=row_idx, column=col_idx).number_format = '@'

    @classmethod
    def delete_dummy_rows(
        cls,
        context: ExcelSheetContext,
        instances: list[OTLObject],
        dummy_data_rows: int,
    ) -> None:
        """Deletes the dummy rows of a template without data rows."""
        if dummy_data_rows != 0:
            return
        instance_count = len([x for x in instances if x.typeURI == context.type_uri])
        if instance_count > 0:
            context.sheet.delete_rows(idx=2, amount=instance_count)

    @classmethod
    def add_deprecated_row(cls, context: ExcelSheetContext) -> None:
        """Adds the row that marks the deprecated attributes of a sheet."""
        if context.info.add_deprecated and any(context.info.deprecated_row):
            cls.add_deprecated_row_to_sheet(context.info.deprecated_row, context.sheet)

    @classmethod
    def add_attribute_info_row(cls, context: ExcelSheetContext) -> None:
        """Adds the row with the definition of the attributes of a sheet."""
        if context.info.add_attribute_info:
            cls.add_attribute_info_to_sheet(context.info.attribute_info_row, context.sheet)

    @classmethod
    def add_deprecated_row_to_sheet(cls, deprecated_attributes_row, sheet):
        sheet.insert_rows(idx=1)
        for index, depr_info in enumerate(deprecated_attributes_row, start=1):
            if depr_info != '':
                cell = sheet.cell(row=1, column=index)
                cell.value = depr_info
                cell.fill = PatternFill(start_color='FF7276', end_color='FF7276', fill_type='solid')

    @classmethod
    def add_attribute_info_to_sheet(cls, collected_attribute_info: list[str], sheet: Worksheet):
        sheet.insert_rows(idx=1)
        fill = PatternFill(start_color='808080', fill_type='solid')
        alignment = Alignment(wrapText=True, vertical='top')
        for index, attr_info in enumerate(collected_attribute_info, start=1):
            cell = sheet.cell(row=1, column=index, value=attr_info)
            cell.alignment = alignment
            cell.fill = fill

    @classmethod
    def set_fixed_column_width(cls, sheet, width: int):
        dim_holder = DimensionHolder(worksheet=sheet)
        for col in range(sheet.min_column, sheet.max_column + 1):
            dim_holder[get_column_letter(col)] = ColumnDimension(
                sheet, min=col, max=col, width=width
            )
        sheet.column_dimensions = dim_holder

    @classmethod
    def abbreviate_excel_sheettitle(cls, sheet):
        class_name = sheet.title
        # abbreviates the class_name so it doesn't exceed the 31 character limit of sheet titles in excel
        if '#' not in class_name:
            sheet.title = class_name[:31]
            return
        split_name = class_name.split('#')
        namespace_name = split_name[0]
        subclass_name = split_name[1]
        class_name = f'{namespace_name[:3]}#{subclass_name}'
        sheet.title = class_name[:31]
