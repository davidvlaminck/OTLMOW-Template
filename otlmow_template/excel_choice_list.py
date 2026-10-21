"""The choice lists of an Excel template.

Every column that holds a boolean or a choice list gets a data validation, and
every choice list gets a column of its own in the 'Keuzelijsten' sheet.
"""

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from otlmow_template.constants import HEADER_EXISTS_AT_COLUMN, MAX_VALIDATION_ROW
from otlmow_template.excel_base import ExcelBaseMixin
from otlmow_template.excel_context import ExcelSheetContext


class ExcelChoiceListMixin(ExcelBaseMixin):
    """The class methods that restrict a column to the values of its choice list."""

    @classmethod
    def add_choice_list_validation(cls, context: ExcelSheetContext, attribute, header_cell) -> None:
        """Adds the validation of a boolean or a choice list column."""
        if not context.generate_choice_list:
            return
        if attribute.field.native_type is bool:
            cls.add_boolean_validation(context=context, header_cell=header_cell)
        elif hasattr(attribute.field, 'options') and attribute.field.options is not None:
            cls.add_option_validation(context=context, attribute=attribute, header_cell=header_cell)

    @classmethod
    def add_boolean_validation(cls, context: ExcelSheetContext, header_cell) -> None:
        """Restricts a boolean column to the values of a boolean list."""
        context.boolean_validation.add(
            cls.get_validation_range(header_cell=header_cell, first_data_row=context.first_data_row)
        )
        cls.color_choice_lists_green(sheet=context.sheet, header_cell_column=header_cell)

    @classmethod
    def add_option_validation(cls, context: ExcelSheetContext, attribute, header_cell) -> None:
        """Restricts a column to the values of its choice list."""
        cls.generate_choice_list_in_excel(
            attribute=attribute,
            choice_list_dict=context.choice_list_dict,
            column=header_cell.column,
            row_nr=1,
            sheet=context.sheet,
            workbook=context.workbook,
        )
        cls.color_choice_lists_green(sheet=context.sheet, header_cell_column=header_cell)

    @classmethod
    def generate_choice_list_in_excel(
        cls, attribute, choice_list_dict, column, row_nr, sheet: Worksheet, workbook: Workbook
    ):
        choice_list_values = [
            cv for cv in attribute.field.options.values() if cv.status != 'verwijderd'
        ]
        if attribute.field.naam not in choice_list_dict:
            cls.add_choice_list_to_sheet(
                workbook=workbook,
                name=attribute.field.naam,
                options=choice_list_values,
                choice_list_dict=choice_list_dict,
            )
        column_in_choice_sheet = choice_list_dict[attribute.field.naam]
        start_range = f'${column_in_choice_sheet}$2'
        end_range = f'${column_in_choice_sheet}${len(choice_list_values) + 2}'
        data_val = DataValidation(
            type='list', formula1=f'Keuzelijsten!{start_range}:{end_range}', allowBlank=True
        )
        sheet.add_data_validation(data_val)
        data_val.add(
            f'{get_column_letter(column)}{row_nr + 1}:'
            f'{get_column_letter(column)}{MAX_VALIDATION_ROW}'
        )

    @classmethod
    def add_choice_list_to_sheet(cls, workbook, name, options, choice_list_dict):
        active_sheet = workbook['Keuzelijsten']
        column_nr = choice_list_dict.keys().__len__() + 1
        row_nr = 1
        new_header = active_sheet.cell(row=row_nr, column=column_nr)
        if new_header.value is not None:
            raise ValueError(
                HEADER_EXISTS_AT_COLUMN.format(column_nr=column_nr, value=new_header.value)
            )
        new_header.value = name
        cell = active_sheet.cell(row=row_nr + 1, column=column_nr)
        cell.value = '-'
        for index, option in enumerate(options, start=2):
            cell = active_sheet.cell(row=row_nr + index, column=column_nr)
            cell.value = option.invulwaarde

        choice_list_dict[name] = new_header.column_letter
