"""The parts of the Excel writer that the sheet writer and the choice lists share."""

from openpyxl.cell import Cell
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from otlmow_converter.Exceptions.UnknownExcelError import UnknownExcelError
from otlmow_model.OtlmowModel.BaseClasses.OTLObject import OTLObject
from otlmow_template.constants import (
    MAX_VALIDATION_ROW,
    SHEET_TITLE_HAS_NO_HASH,
    short_to_long_ns,
)


class ExcelBaseMixin:
    """The class methods that every part of the Excel writer needs."""

    green_fill = PatternFill(start_color='90EE90', fill_type='solid')

    @classmethod
    def get_validation_range(cls, header_cell, first_data_row: int) -> str:
        """Returns the range of the rows a validation is added to."""
        column_letter = header_cell.column_letter
        return f'{column_letter}{first_data_row}:{column_letter}{MAX_VALIDATION_ROW}'

    @classmethod
    def color_choice_lists_green(cls, sheet: Worksheet, header_cell_column: Cell):
        for cell in sheet.iter_rows(
            min_col=header_cell_column.column,
            max_col=header_cell_column.column,
            min_row=2,
            max_row=MAX_VALIDATION_ROW,
        ):
            cell[0].fill = cls.green_fill

    @classmethod
    def get_instance_for_type_uri(cls, instances: list[OTLObject], type_uri: str):
        """Returns the instance of the type of a sheet, raises when there is none."""
        instance = cls.find_instance_with_type_uri(instances=instances, type_uri=type_uri)
        if instance is None:
            instance = cls.find_instance_with_type_uri_prefix(
                instances=instances, type_uri=type_uri
            )
        if instance is None:
            raise UnknownExcelError(
                f'When creating a template, no instance could be created for {type_uri}'
            )
        return instance

    @classmethod
    def find_instance_with_type_uri(cls, instances: list[OTLObject], type_uri: str):
        """Returns the instance with exactly the given type uri, None when there is none."""
        for instance in instances:
            if instance.typeURI == type_uri:
                return instance
        return None

    @classmethod
    def find_instance_with_type_uri_prefix(cls, instances: list[OTLObject], type_uri: str):
        """Returns the instance whose type uri starts with the given one, None when there is none."""
        for instance in instances:
            if instance.typeURI.startswith(type_uri):
                return instance
        return None

    @classmethod
    def get_first_data_row(cls, add_attribute_info: bool) -> int:
        """Returns the first row of a sheet that holds data.

        The first row is the header row. The attribute info row is added in front
        of it, so the data starts one row lower. The deprecated row is added
        after the validations are applied, so it does not move the data.
        """
        first_data_row = 2
        if add_attribute_info:
            first_data_row += 1
        return first_data_row

    @classmethod
    def get_uri_from_sheet_name(cls, title: str) -> str:
        if title == 'Agent':
            return 'http://purl.org/dc/terms/Agent'
        if '#' not in title:
            raise ValueError(SHEET_TITLE_HAS_NO_HASH)
        class_ns, class_name = title.split('#', maxsplit=1)
        return short_to_long_ns.get(class_ns, class_ns) + class_name
