import asyncio
import concurrent
import contextlib
import csv
import logging
from asyncio import sleep
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed, ALL_COMPLETED
from pathlib import Path

from openpyxl.cell import Cell
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import DimensionHolder, ColumnDimension
from openpyxl.worksheet.worksheet import Worksheet
from otlmow_converter.DotnotationHelper import DotnotationHelper
from otlmow_converter.Exceptions.UnknownExcelError import UnknownExcelError
from otlmow_converter.OtlmowConverter import OtlmowConverter
from otlmow_model.OtlmowModel.BaseClasses.OTLObject import dynamic_create_instance_from_uri, OTLObject, \
    get_attribute_by_name, dynamic_create_type_from_uri, OTLAttribuut
from otlmow_model.OtlmowModel.Helpers.GenericHelper import get_ns_and_name_from_uri
from otlmow_model.OtlmowModel.Helpers.RelationCreator import create_betrokkenerelation, create_relation
from otlmow_model.OtlmowModel.Helpers.generated_lists import get_hardcoded_relation_dict
from otlmow_modelbuilder.OSLOCollector import OSLOCollector
from otlmow_modelbuilder.SQLDataClasses.OSLOClass import OSLOClass

from otlmow_template.TemplateGeneratorBase import TemplateGeneratorBase

short_to_long_ns = {
    'ond': 'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#',
    'onderdeel': 'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#',
    'ins': 'https://wegenenverkeer.data.vlaanderen.be/ns/installatie#',
    'installatie': 'https://wegenenverkeer.data.vlaanderen.be/ns/installatie#',
    'imp': 'https://wegenenverkeer.data.vlaanderen.be/ns/implementatieelement#',
    'implementatieelement': 'https://wegenenverkeer.data.vlaanderen.be/ns/implementatieelement#',
    'proefenmeting': 'https://wegenenverkeer.data.vlaanderen.be/ns/proefenmeting#',
    'pro': 'https://wegenenverkeer.data.vlaanderen.be/ns/proefenmeting#',
    'lev': 'https://wegenenverkeer.data.vlaanderen.be/ns/levenscyclus#',
    'levenscyclus': 'https://wegenenverkeer.data.vlaanderen.be/ns/levenscyclus#'
}


class AanlevertemplateGenerator(TemplateGeneratorBase):
    """Generates a delivery template (aanlevertemplate) from a subset database.
    
    The delivery template is used to provide data according to the OTL standard.
    It contains dummy data rows and choice lists for validation.
    """
    
    green_fill = PatternFill(start_color="90EE90", fill_type="solid")
    
    def __init__(self):
        super().__init__()
        self.relation_dict = {}
    
    @classmethod
    def filters_classes_by_subset(cls, collector: OSLOCollector,
                                  class_uris_filter: [str] = None) -> list[OSLOClass]:
        if class_uris_filter is None:
            return collector.classes
        return [x for x in collector.classes if x.objectUri in class_uris_filter]
    
    def create_x_objects(self, oslo_class, add_geometry, collector, filter_attributes_by_subset,
                         model_directory, amount_objects_to_create):
        if oslo_class.objectUri in self.relation_dict:
            return []
        
        otl_objects = []
        for _ in range(amount_objects_to_create):
            otl_object = self.generate_object_from_oslo_class(
                oslo_class=oslo_class, add_geometry=add_geometry, collector=collector,
                filter_attributes_by_subset=filter_attributes_by_subset, model_directory=model_directory)
            if otl_object is not None:
                otl_objects.append(otl_object)
        return otl_objects
    
    def generate_objects_for_template(
            self, subset_path: Path, class_uris_filter: [str], filter_attributes_by_subset: bool,
            dummy_data_rows: int, add_geometry: bool, ignore_relations: bool, model_directory: Path = None
    ) -> [OTLObject]:
        """
        This method is used to generate objects for the template. It will generate objects based on the subset file
        """
        collector = self._load_collector_from_subset_path(subset_path=subset_path)
        filtered_class_list = self.filters_classes_by_subset(collector=collector, class_uris_filter=class_uris_filter)
        if filtered_class_list == []:
            raise ValueError('Something went wrong, as the class_uri filter list is empty')
        self.relation_dict = get_hardcoded_relation_dict(model_directory=model_directory)
        
        amount_objects_to_create = max(1, dummy_data_rows)
        otl_objects = []
        
        while True:
            with ThreadPoolExecutor() as executor:
                futures = [
                    executor.submit(self.create_x_objects, oslo_class, add_geometry, collector,
                                    filter_attributes_by_subset, model_directory, amount_objects_to_create)
                    for oslo_class in [cl for cl in filtered_class_list if cl.abstract == 0]
                ]
                while futures:
                    done, not_done = concurrent.futures.wait(futures, return_when=ALL_COMPLETED, timeout=60)
                    for future in as_completed(done):
                        otl_objects.extend(future.result())
                    futures = not_done
                    print(f'when using multiprocessing: {len(done)} done, {len(not_done)} not done')
                    if len(not_done) > 0:
                        print('restarting the loop with the not done ones')
            
            created = len(otl_objects)
            unique_ids = len({obj.assetId.identificator if obj.typeURI != 'http://purl.org/dc/terms/Agent' else obj.agentId.identificator
                              for obj in otl_objects})
            if created == unique_ids:
                break
            otl_objects = []
        
        if not ignore_relations:
            non_relations_class_uris = [cl.objectUri for cl in filtered_class_list
                                        if cl.abstract == 0 and cl.objectUri not in self.relation_dict]
            self.append_relations_to_objects(otl_objects=otl_objects, collector=collector,
                                            class_uris_filter=non_relations_class_uris, model_directory=model_directory)
        
        return otl_objects
    
    async def create_x_objects_async(self, oslo_class, add_geometry, collector, filter_attributes_by_subset,
                                  model_directory, amount_objects_to_create):
        if oslo_class.objectUri in self.relation_dict:
            return []
        
        otl_objects = []
        for _ in range(amount_objects_to_create):
            await sleep(0)
            otl_object = self.generate_object_from_oslo_class(
                oslo_class=oslo_class, add_geometry=add_geometry, collector=collector,
                filter_attributes_by_subset=filter_attributes_by_subset, model_directory=model_directory)
            if otl_object is not None:
                otl_objects.append(otl_object)
        return otl_objects
    
    async def generate_objects_for_template_async(
            self, subset_path: Path, class_uris_filter: [str], filter_attributes_by_subset: bool,
            dummy_data_rows: int, add_geometry: bool, ignore_relations: bool, model_directory: Path = None
    ) -> [OTLObject]:
        """
        This method is used to generate objects for the template. It will generate objects based on the subset file
        """
        await sleep(0)
        collector = self._load_collector_from_subset_path(subset_path=subset_path)
        await sleep(0)
        filtered_class_list = self.filters_classes_by_subset(collector=collector, class_uris_filter=class_uris_filter)
        if filtered_class_list == []:
            raise ValueError('Something went wrong, as the class_uri filter list is empty')
        await sleep(0)
        self.relation_dict = get_hardcoded_relation_dict(model_directory=model_directory)
        
        amount_objects_to_create = max(1, dummy_data_rows)
        otl_objects = []
        
        while True:
            tasks = [
                self.create_x_objects_async(oslo_class, add_geometry, collector,
                                           filter_attributes_by_subset, model_directory, amount_objects_to_create)
                for oslo_class in [cl for cl in filtered_class_list if cl.abstract == 0]
            ]
            results = await asyncio.gather(*tasks)
            otl_objects.extend([item for sublist in results for item in sublist])
            
            created = len(otl_objects)
            unique_ids = len({
                                 obj.assetId.identificator if obj.typeURI != 'http://purl.org/dc/terms/Agent' else obj.agentId.identificator
                                 for obj in otl_objects})
            if created == unique_ids:
                break
            otl_objects = []
        
        if not ignore_relations:
            non_relations_class_uris = [cl.objectUri for cl in filtered_class_list
                                        if cl.abstract == 0 and cl.objectUri not in self.relation_dict]
            self.append_relations_to_objects(otl_objects=otl_objects, collector=collector,
                                            class_uris_filter=non_relations_class_uris, model_directory=model_directory)
        
        return otl_objects
    
    def generate_object_from_oslo_class(
            self, oslo_class: OSLOClass, add_geometry: bool,
            filter_attributes_by_subset: bool, collector: OSLOCollector, model_directory: Path = None) -> [OTLObject]:
        """
        Generate an object from a given OSLO class
        """
        instance = dynamic_create_instance_from_uri(oslo_class.objectUri, model_directory=model_directory)
        if instance is None:
            return
        
        if filter_attributes_by_subset:
            for attribute_object in collector.find_attributes_by_class(oslo_class):
                attr = get_attribute_by_name(instance, attribute_object.name)
                if attr is None:
                    logging.warning(f'Attribute {attribute_object.name} not found in class {oslo_class.objectUri}')
                elif attr.naam == 'isActief':
                    attr.set_waarde(True)
                else:
                    attr.fill_with_dummy_data()
        else:
            for attr in instance:
                if attr.naam == 'isActief':
                    attr.set_waarde(True)
                elif attr.naam != 'geometry':
                    attr.fill_with_dummy_data()
        
        with contextlib.suppress(AttributeError):
            assetId_attr = get_attribute_by_name(instance, 'assetId')
            if assetId_attr is not None:
                assetId_attr.fill_with_dummy_data()
            if add_geometry:
                geo_attr = get_attribute_by_name(instance, 'geometry')
                if geo_attr is not None:
                    geo_attr.fill_with_dummy_data()
        
        asset_versie = get_attribute_by_name(instance, 'assetVersie')
        if asset_versie is not None:
            asset_versie.set_waarde(None)
        
        DotnotationHelper.clear_list_of_list_attributes(instance)
        
        return instance
    
    def alter_excel_template(self, instances: list, file_path: Path, add_attribute_info: bool,
                             generate_choice_list: bool, dummy_data_rows: int, add_deprecated: bool,
                             abbreviate_excel_sheettitles: bool):
        wb = load_workbook(file_path)
        wb.create_sheet('Keuzelijsten')
        
        choice_list_dict = {}
        for sheet in wb:
            if sheet.title == 'Keuzelijsten':
                break
            
            self.alter_excel_sheet(add_attribute_info=add_attribute_info, choice_list_dict=choice_list_dict,
                                  generate_choice_list=generate_choice_list, dummy_data_rows=dummy_data_rows,
                                  instances=instances, sheet=sheet, add_deprecated=add_deprecated, workbook=wb,
                                  abbreviate_excel_sheettitle=abbreviate_excel_sheettitles)
        
        wb.save(file_path)
        wb.close()
    
    def fill_class_dict(self, instances: list) -> dict:
        class_dict = defaultdict(list)
        for instance in instances:
            class_dict[instance.typeURI].append(instance)
        return class_dict
    
    def alter_csv_template(self, instances: list, file_path: Path, add_attribute_info: bool,
                           split_per_type: bool, dummy_data_rows: int, add_deprecated: bool):
        classes_dict = self.fill_class_dict(instances)
        if split_per_type:
            for type_uri, typed_instances in classes_dict.items():
                ns, name = get_ns_and_name_from_uri(type_uri)
                ns_name = f'{ns}_{name}'
                if name == 'Agent':
                    ns_name = 'Agent'
                class_file_path = file_path.parent / f'{file_path.stem}_{ns_name}.csv'
                self.alter_csv_file(add_attribute_info=add_attribute_info, add_deprecated=add_deprecated,
                                   dummy_data_rows=dummy_data_rows, instances=typed_instances, file_path=class_file_path)
        else:
            self.alter_csv_file(add_attribute_info=add_attribute_info, add_deprecated=add_deprecated,
                               dummy_data_rows=dummy_data_rows, instances=instances, file_path=file_path)
    
    async def alter_csv_template_async(self, instances: list, file_path: Path, add_deprecated: bool,
                                       add_attribute_info: bool, split_per_type: bool, dummy_data_rows: int):
        classes_dict = self.fill_class_dict(instances)
        if split_per_type:
            for type_uri, typed_instances in classes_dict.items():
                await sleep(0)
                ns, name = get_ns_and_name_from_uri(type_uri)
                ns_name = f'{ns}_{name}'
                if name == 'Agent':
                    ns_name = 'Agent'
                class_file_path = file_path.parent / f'{file_path.stem}_{ns_name}.csv'
                self.alter_csv_file(add_attribute_info=add_attribute_info,
                                   dummy_data_rows=dummy_data_rows, instances=typed_instances, add_deprecated=add_deprecated,
                                   file_path=class_file_path)
        else:
            await sleep(0)
            self.alter_csv_file(add_attribute_info=add_attribute_info,
                               dummy_data_rows=dummy_data_rows, instances=instances, add_deprecated=add_deprecated,
                               file_path=file_path)
    
    def alter_csv_file(self, add_attribute_info: bool, instances: [OTLObject], add_deprecated: bool, file_path: Path,
                       dummy_data_rows: int):
        collected_attribute_info_row = []
        deprecated_attributes_row = []
        instance = instances[0]
        quote_char = '"'
        
        with open(file_path, encoding='utf-8') as file:
            csv_reader = csv.reader(file, delimiter=';', quotechar=quote_char)
            header_row = next(csv_reader)
            csv_data = list(csv_reader)
        
        for index, header in enumerate(header_row):
            if header is None or header == '':
                continue
            
            if header == 'typeURI':
                if add_attribute_info:
                    collected_attribute_info_row.append(
                        'De URI van het object volgens https://www.w3.org/2001/XMLSchema#anyURI .')
                if add_deprecated:
                    deprecated_attributes_row.append('')
                continue
            
            try:
                attribute = DotnotationHelper.get_attribute_by_dotnotation(instance, header)
            except AttributeError:
                if add_attribute_info:
                    collected_attribute_info_row.append('')
                if add_deprecated:
                    deprecated_attributes_row.append('')
                continue
            
            if add_attribute_info:
                collected_attribute_info_row.append(
                    self.get_attribute_info(instance=instance, header=header, attribute=attribute))
            
            if add_deprecated:
                deprecated_attributes_row.append('DEPRECATED' if attribute.deprecated_version else '')
        
        with open(file_path, 'w', newline='', encoding='utf-8') as file:
            csv_writer = csv.writer(file, delimiter=';', quotechar=quote_char, quoting=csv.QUOTE_MINIMAL)
            if add_attribute_info:
                csv_writer.writerow(collected_attribute_info_row)
            if add_deprecated and any(deprecated_attributes_row):
                csv_writer.writerow(deprecated_attributes_row)
            csv_writer.writerow(header_row)
            if dummy_data_rows != 0:
                for line in csv_data:
                    csv_writer.writerow(line)
    
    async def alter_excel_template_async(self, instances: list, file_path: Path, add_attribute_info: bool,
                                          generate_choice_list: bool, dummy_data_rows: int, add_deprecated: bool,
                                          abbreviate_excel_sheettitles: bool):
        wb = load_workbook(file_path)
        wb.create_sheet('Keuzelijsten')
        
        choice_list_dict = {}
        for sheet in wb:
            if sheet.title == 'Keuzelijsten':
                break
            
            self.alter_excel_sheet(add_attribute_info=add_attribute_info, choice_list_dict=choice_list_dict,
                                  generate_choice_list=generate_choice_list, dummy_data_rows=dummy_data_rows,
                                  instances=instances, sheet=sheet, add_deprecated=add_deprecated, workbook=wb,
                                  abbreviate_excel_sheettitle=abbreviate_excel_sheettitles)
            await sleep(0)
        
        wb.save(file_path)
        wb.close()
    
    def alter_excel_sheet(self, add_attribute_info: bool, choice_list_dict: dict, generate_choice_list: bool,
                          instances: [OTLObject], sheet: Worksheet, add_deprecated: bool, workbook: Workbook,
                          dummy_data_rows: int, abbreviate_excel_sheettitle: bool = True):
        type_uri = self.get_uri_from_sheet_name(sheet.title)
        instance = next((x for x in instances if x.typeURI == type_uri), None)
        if instance is None:
            instance = next((x for x in instances if x.typeURI.startswith(type_uri)), None)
        if instance is None:
            raise UnknownExcelError(f'When creating a template, no instance could be created for {type_uri}')
        
        boolean_validation = DataValidation(type="list", formula1='"TRUE,FALSE,-"', allow_blank=True)
        sheet.add_data_validation(boolean_validation)
        collected_attribute_info = []
        deprecated_attributes_row = []
        
        header_row = next(sheet.iter_rows(min_row=1, max_row=1))
        
        first_data_row = 2
        if add_attribute_info:
            first_data_row += 1
        if add_deprecated and any(deprecated_attributes_row):
            first_data_row += 1
        
        max_row = 1000
        string_col_indices = []
        
        for header_cell in header_row:
            header = header_cell.value
            if not header:
                continue
            
            if header == 'typeURI':
                data_validation = DataValidation(type="list", formula1=f'"{type_uri}"', allow_blank=True)
                sheet.add_data_validation(data_validation)
                data_validation.add(f'{header_cell.column_letter}{first_data_row}:{header_cell.column_letter}{max_row}')
                if add_attribute_info:
                    collected_attribute_info.append(
                        'De URI van het object volgens https://www.w3.org/2001/XMLSchema#anyURI .')
                if add_deprecated:
                    deprecated_attributes_row.append('')
                continue
            
            if type_uri == 'http://purl.org/dc/terms/Agent' and header.startswith('assetId.'):
                continue
            
            attribute = DotnotationHelper.get_attribute_by_dotnotation(instance, header)
            
            if add_attribute_info:
                collected_attribute_info.append(
                    self.get_attribute_info(instance=instance, header=header, attribute=attribute))
            
            if add_deprecated:
                deprecated_attributes_row.append('DEPRECATED' if attribute.deprecated_version else '')
            
            if generate_choice_list:
                if attribute.field.native_type == bool:
                    boolean_validation.add(f'{header_cell.column_letter}{first_data_row}:{header_cell.column_letter}{max_row}')
                    self.color_choice_lists_green(sheet=sheet, header_cell_column=header_cell)
                elif hasattr(attribute.field, 'options') and attribute.field.options is not None:
                    self.generate_choice_list_in_excel(
                        attribute=attribute, choice_list_dict=choice_list_dict, column=header_cell.column,
                        row_nr=1, sheet=sheet, workbook=workbook)
                    self.color_choice_lists_green(sheet=sheet, header_cell_column=header_cell)
            
            if attribute.field.native_type == str:
                string_col_indices.append(header_cell.column)
        
        for col_idx in string_col_indices:
            for row_idx in range(first_data_row, max_row + 1):
                sheet.cell(row=row_idx, column=col_idx).number_format = "@"
        
        if dummy_data_rows == 0:
            instance_count = len([x for x in instances if x.typeURI == type_uri])
            if instance_count > 0:
                sheet.delete_rows(idx=2, amount=instance_count)
        
        if add_deprecated and any(deprecated_attributes_row):
            self.add_deprecated_row_to_sheet(deprecated_attributes_row, sheet)
        
        if add_attribute_info:
            self.add_attribute_info_to_sheet(collected_attribute_info, sheet)
        
        self.set_fixed_column_width(sheet=sheet, width=25)
        
        if abbreviate_excel_sheettitle:
            self.abbreviate_excel_sheettitle(sheet=sheet)
    
    def color_choice_lists_green(self, sheet: Worksheet, header_cell_column: Cell):
        for cell in sheet.iter_rows(min_col=header_cell_column.column, max_col=header_cell_column.column,
                                    min_row=2, max_row=1000):
            cell[0].fill = self.green_fill
    
    def add_deprecated_row_to_sheet(self, deprecated_attributes_row, sheet):
        sheet.insert_rows(idx=1)
        for index, depr_info in enumerate(deprecated_attributes_row, start=1):
            if depr_info != '':
                cell = sheet.cell(row=1, column=index)
                cell.value = depr_info
                cell.fill = PatternFill(start_color="FF7276", end_color="FF7276", fill_type="solid")
    
    def add_attribute_info_to_sheet(self, collected_attribute_info: list[str], sheet: Worksheet):
        sheet.insert_rows(idx=1)
        fill = PatternFill(start_color="808080", fill_type="solid")
        alignment = Alignment(wrapText=True, vertical='top')
        for index, attr_info in enumerate(collected_attribute_info, start=1):
            cell = sheet.cell(row=1, column=index, value=attr_info)
            cell.alignment = alignment
            cell.fill = fill
    
    def get_attribute_info(self, instance: OTLObject, header: str, attribute: OTLAttribuut) -> str:
        """Returns the definition to show in the attribute info row for a column of a template."""
        quantitative_attribute = self.get_quantitative_attribute(instance=instance, header=header)
        if quantitative_attribute is None:
            return attribute.definition
        standard_unit = self.get_standard_unit(quantitative_attribute)
        if standard_unit is None:
            return attribute.definition
        definition = quantitative_attribute.definition.strip()
        if not definition.endswith('.'):
            definition = f'{definition}.'
        return f'{definition} Standaard eenheid: {standard_unit}'
    
    def get_quantitative_attribute(self, instance: OTLObject, header: str) -> OTLAttribuut:
        """Returns the attribute a column header refers to when it points to the value of a quantitative value."""
        attribute = DotnotationHelper.get_attribute_by_dotnotation(
            instance_or_attribute=instance, dotnotation=header, waarde_shortcut=False)
        if not attribute.field.waarde_shortcut_applicable:
            return None
        if not hasattr(attribute.field.waardeObject, 'standaardEenheid'):
            return None
        return attribute
    
    def get_standard_unit(self, attribute: OTLAttribuut) -> str:
        """Returns the standard unit (ex. 'mo') of a quantitative value attribute."""
        waarde = attribute.waarde
        if waarde is None:
            attribute.add_empty_value()
            waarde = attribute.waarde
        if isinstance(waarde, list):
            waarde = next(iter(waarde), None)
        standaard_eenheid = get_attribute_by_name(waarde, 'standaardEenheid')
        if standaard_eenheid is None:
            return None
        for notation in [standaard_eenheid.usagenote, standaard_eenheid.constraints]:
            if '"' in notation:
                return notation.split('"')[1]
        return None
    
    def generate_choice_list_in_excel(self, attribute, choice_list_dict, column, row_nr, sheet: Worksheet,
                                      workbook: Workbook):
        choice_list_values = [cv for cv in attribute.field.options.values()
                              if cv.status != 'verwijderd']
        if attribute.field.naam not in choice_list_dict:
            self.add_choice_list_to_sheet(workbook=workbook, name=attribute.field.naam,
                                         options=choice_list_values, choice_list_dict=choice_list_dict)
        column_in_choice_sheet = choice_list_dict[attribute.field.naam]
        start_range = f"${column_in_choice_sheet}$2"
        end_range = f"${column_in_choice_sheet}${len(choice_list_values) + 2}"
        data_val = DataValidation(type="list", formula1=f"Keuzelijsten!{start_range}:{end_range}",
                                  allowBlank=True)
        sheet.add_data_validation(data_val)
        data_val.add(f'{get_column_letter(column)}{row_nr + 1}:'
                     f'{get_column_letter(column)}1000')
    
    def set_fixed_column_width(self, sheet, width: int):
        dim_holder = DimensionHolder(worksheet=sheet)
        for col in range(sheet.min_column, sheet.max_column + 1):
            dim_holder[get_column_letter(col)] = ColumnDimension(sheet, min=col, max=col, width=width)
        sheet.column_dimensions = dim_holder
    
    def add_choice_list_to_sheet(self, workbook, name, options, choice_list_dict):
        active_sheet = workbook['Keuzelijsten']
        column_nr = choice_list_dict.keys().__len__() + 1
        row_nr = 1
        new_header = active_sheet.cell(row=row_nr, column=column_nr)
        if new_header.value is not None:
            raise ValueError(f'Header already exists at column {column_nr}: {new_header.value}')
        new_header.value = name
        cell = active_sheet.cell(row=row_nr + 1, column=column_nr)
        cell.value = "-"
        for index, option in enumerate(options, start=2):
            cell = active_sheet.cell(row=row_nr + index, column=column_nr)
            cell.value = option.invulwaarde
        
        choice_list_dict[name] = new_header.column_letter
    
    def get_uri_from_sheet_name(self, title: str) -> str:
        if title == 'Agent':
            return 'http://purl.org/dc/terms/Agent'
        if '#' not in title:
            raise ValueError('Sheet title does not contain a #')
        class_ns, class_name = title.split('#', maxsplit=1)
        return short_to_long_ns.get(class_ns, class_ns) + class_name
    
    def append_relations_to_objects(self, otl_objects: [OTLObject], collector: OSLOCollector, class_uris_filter: [str],
                                    model_directory: Path):
        class_dict = defaultdict(list)
        for instance in otl_objects:
            class_dict[instance.typeURI].append(instance)
        
        for class_uri in class_uris_filter:
            for relation in collector.find_all_concrete_relations(objectUri=class_uri):
                if class_uri != relation.bron_uri:
                    continue
                for i, bron_instance in enumerate(class_dict[relation.bron_uri]):
                    if relation.doel_uri not in class_dict:
                        continue
                    doel_instance = class_dict[relation.doel_uri][i]
                    if relation.objectUri == 'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#HeeftBetrokkene':
                        relation_instance = create_betrokkenerelation(rol='toezichter', source=bron_instance,
                                                                    target=doel_instance, model_directory=model_directory)
                    else:
                        if relation.richting == 'Unspecified' and bron_instance.assetId.identificator > doel_instance.assetId.identificator:
                            continue
                        relation_type = dynamic_create_type_from_uri(class_uri=relation.objectUri, model_directory=model_directory)
                        relation_instance = create_relation(relation_type=relation_type, source=bron_instance,
                                                            target=doel_instance, model_directory=model_directory)
                    
                    otl_objects.append(relation_instance)
    
    def abbreviate_excel_sheettitle(self, sheet):
        class_name = sheet.title
        if '#' not in class_name:
            sheet.title = class_name[:31]
            return
        split_name = class_name.split("#")
        namespace_name = split_name[0]
        subclass_name = split_name[1]
        class_name = f"{namespace_name[:3]}#{subclass_name}"
        sheet.title = class_name[:31]