"""The guard and error paths of SubsetTemplateCreator.

These are the paths a caller reaches when something is wrong: an empty filter, a
sheet title without a namespace, a header that is already in use, an attribute
that does not exist. The happy-path suite never reaches them, so a changed guard
or a changed error message is invisible.

The paths that *swallow* a problem are the most important ones: they are the
places where a defect turns into a template that looks fine but is not.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from otlmow_converter.Exceptions.UnknownExcelError import UnknownExcelError
from otlmow_template.SubsetTemplateCreator import SubsetTemplateCreator

current_dir = Path(__file__).parent
model_directory_path = Path(__file__).parent.parent / 'TestModel'


def test_get_uri_from_sheet_name_requires_a_namespace():
    """A sheet title without a '#' cannot be mapped to a class."""
    with pytest.raises(ValueError, match='does not contain a #'):
        SubsetTemplateCreator.get_uri_from_sheet_name('GeenNamespace')


def test_get_uri_from_sheet_name_maps_the_agent_sheet():
    """The Agent sheet has its own uri."""
    assert (
        SubsetTemplateCreator.get_uri_from_sheet_name('Agent') == 'http://purl.org/dc/terms/Agent'
    )


def test_get_uri_from_sheet_name_expands_the_short_namespace():
    """A short namespace is expanded to the full uri."""
    assert (
        SubsetTemplateCreator.get_uri_from_sheet_name('onderdeel#AllCasesTestClass')
        == 'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#AllCasesTestClass'
    )


def test_filters_classes_by_subset_returns_nothing_for_an_unknown_filter():
    """A filter that matches no class gives an empty list."""
    collector = SubsetTemplateCreator._load_collector_from_subset_path(
        subset_path=current_dir / 'OTL_AllCasesTestClass.db'
    )
    assert (
        SubsetTemplateCreator.filters_classes_by_subset(
            collector=collector, class_uris_filter=['https://example.com/does-not-exist']
        )
        == []
    )


def test_an_empty_class_uris_filter_raises(tmp_path):
    """Generating a template with a filter that matches nothing is an error."""
    with pytest.raises(ValueError, match='class_uri filter list is empty'):
        SubsetTemplateCreator().generate_template_from_subset(
            subset_path=current_dir / 'OTL_AllCasesTestClass.db',
            template_file_path=tmp_path / 'template.xlsx',
            class_uris_filter=['https://example.com/does-not-exist'],
            model_directory=model_directory_path,
        )


@pytest.mark.asyncio
async def test_an_empty_class_uris_filter_raises_async(tmp_path):
    """The asynchronous variant must raise the same error."""
    with pytest.raises(ValueError, match='class_uri filter list is empty'):
        await SubsetTemplateCreator().generate_template_from_subset_async(
            subset_path=current_dir / 'OTL_AllCasesTestClass.db',
            template_file_path=tmp_path / 'template.xlsx',
            class_uris_filter=['https://example.com/does-not-exist'],
            model_directory=model_directory_path,
        )


def test_a_duplicate_choice_list_column_is_refused():
    """A choice list cannot overwrite an existing column."""
    workbook = Workbook()
    workbook.create_sheet('Keuzelijsten')
    workbook['Keuzelijsten'].cell(row=1, column=3).value = 'BestaandeKeuze'
    choice_list_dict = {'eerste': 'A', 'tweede': 'B'}

    with pytest.raises(ValueError, match='Header already exists'):
        SubsetTemplateCreator.add_choice_list_to_sheet(
            workbook=workbook,
            name='derde',
            options=[],
            choice_list_dict=choice_list_dict,
        )


def test_an_unknown_attribute_is_skipped_in_the_csv(tmp_path):
    """An attribute that the instance does not have is left empty."""
    instances = SubsetTemplateCreator.generate_objects_for_template(
        subset_path=current_dir / 'OTL_AllCasesTestClass.db',
        class_uris_filter=[
            'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#AnotherTestClass'
        ],
        filter_attributes_by_subset=True,
        dummy_data_rows=0,
        add_geometry=True,
        ignore_relations=True,
        model_directory=model_directory_path,
    )
    assert instances

    csv_path = tmp_path / 'template.csv'
    csv_path.write_text(
        'typeURI;assetId.identificator;bestaatNiet\n'
        'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#AnotherTestClass;dummy;dummy\n',
        encoding='utf-8',
    )

    SubsetTemplateCreator.alter_csv_file(
        add_attribute_info=True,
        instances=instances,
        add_deprecated=True,
        file_path=csv_path,
        dummy_data_rows=0,
    )

    rows = csv_path.read_text(encoding='utf-8').splitlines()
    # the attribute information row is added in front of the header row; the
    # deprecated row is only added when at least one attribute is deprecated
    assert len(rows) == 2
    header = rows[1].split(';')
    information = rows[0].split(';')
    unknown_index = header.index('bestaatNiet')
    assert information[unknown_index] == ''
    assert information[header.index('assetId.identificator')] != ''


def test_a_sheet_without_a_matching_instance_raises():
    """A sheet that has no matching instance is an error, not an empty sheet."""
    instances = SubsetTemplateCreator.generate_objects_for_template(
        subset_path=current_dir / 'OTL_AllCasesTestClass.db',
        class_uris_filter=[
            'https://wegenenverkeer.data.vlaanderen.be/ns/onderdeel#AnotherTestClass'
        ],
        filter_attributes_by_subset=True,
        dummy_data_rows=0,
        add_geometry=True,
        ignore_relations=True,
        model_directory=model_directory_path,
    )
    assert instances

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'onderdeel#BestaatNiet'

    with pytest.raises(UnknownExcelError, match='no instance could be created'):
        SubsetTemplateCreator.alter_excel_sheet(
            add_attribute_info=False,
            choice_list_dict={},
            generate_choice_list=False,
            dummy_data_rows=0,
            instances=instances,
            sheet=sheet,
            add_deprecated=False,
            workbook=workbook,
        )
